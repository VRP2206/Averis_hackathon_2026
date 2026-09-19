"""Build a pack of test .eml files, one per situation SDOC must handle, then
check each one against the pipeline.  Run from the repo root:

    python scripts/make_test_emails.py            # writes docs/test-emails/*.eml
"""
from __future__ import annotations

import io
import sys
from email.message import EmailMessage
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "test-emails"
sys.path.insert(0, str(ROOT))

SIG = "\n\nBest Regards,\nAisyah Rahman\nShipping Documentation\nDID : +60 3 2710 5521\nDemo Trading Sdn Bhd"

# One canonical shipment; the BL copies it with deliberate changes per case.
SHIP = {
    "shipper": "DEMO PAPER TRADING SDN BHD",
    "consignee": "HARBOURLINE STATIONERY LLC",
    "notify_party": "NORTHWIND OFFICE SUPPLIES PTE LTD",
    "port_of_loading": "PORT KLANG (WESTPORT), MALAYSIA (MYPKG)",
    "port_of_discharge": "JEBEL ALI, UNITED ARAB EMIRATES (AEJEA)",
    "container_count": "4 x 40'HC",
    "gross_weight_kg": "86,400 KG",
}
SI_LABELS = {"shipper": "Shipper/Exporter", "consignee": "Consignee", "notify_party": "Notify Party",
             "port_of_loading": "Port of Loading", "port_of_discharge": "Discharge Port",
             "container_count": "No. of Containers", "gross_weight_kg": "Gross Weight (KG)"}
BL_LABELS = {"shipper": "SHIPPER", "consignee": "To the Order of", "notify_party": "Notify",
             "port_of_loading": "POL", "port_of_discharge": "Port of Discharge (POD)",
             "container_count": "Container Count", "gross_weight_kg": "Gross Wt (kgs)"}


def si_txt(values=SHIP) -> bytes:
    lines = ["SHIPPING INSTRUCTION", "=" * 40, ""] + [f"{SI_LABELS[k]}: {v}" for k, v in values.items()]
    return ("\n".join(lines + ["Commodity: A4 COPIER PAPER 80GSM", "Booking Ref: DEMO-BK-2026-0917"]) + "\n").encode()


def bl_txt(values=SHIP) -> bytes:
    lines = ["BILL OF LADING (DRAFT)", "=" * 40, ""] + [f"{BL_LABELS[k]}: {v}" for k, v in values.items()]
    return ("\n".join(lines + ["Vessel Name: DEMO STAR V.026E", "Bill of Lading No.: DEMOBL260917"]) + "\n").encode()


def si_xlsx(values=SHIP) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "S.I."
    ws.append(["BL INSTRUCTION", "SO-0917"])
    for k, v in values.items():
        ws.append([SI_LABELS[k], v])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def bl_docx(values=SHIP) -> bytes:
    from docx import Document
    d = Document(); d.add_heading("BILL OF LADING (DRAFT)", level=1)
    t = d.add_table(rows=0, cols=2)
    for k, v in values.items():
        cells = t.add_row().cells; cells[0].text = BL_LABELS[k]; cells[1].text = v
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def invoice_txt() -> bytes:
    return (b"COMMERCIAL INVOICE\n========================================\n\nInvoice No.: CI-2026-0917\n"
            b"Seller: DEMO PAPER TRADING SDN BHD\nBuyer: HARBOURLINE STATIONERY LLC\nTotal: USD 48,960.00\n")


def email(subject: str, body: str, sender="aisyah@demotrading.example", files: list[tuple[str, bytes]] = ()) -> bytes:
    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = sender, "docs@sdoc-demo.example", subject
    m.set_content(body)
    for name, data in files:
        m.add_attachment(data, maintype="application", subtype="octet-stream", filename=name)
    return m.as_bytes()


def with_(**changes):
    return {**SHIP, **changes}


CASES = [
    # file, expected category, expected status, expected fields/reason, bytes
    ("01-mismatch-port-and-containers.eml", "BL_COMPARISON", "MISMATCH", "container_count, port_of_discharge",
     email("TO CONFIRM DOCS _ 5RDM-00917 _ JEBEL ALI_UAE _ HARBOURLINE STATIONERY LLC _ DEMOBL260917",
           "Hi team,\n\nAttached are the SI and draft BL for OC 5RDM-00917. Please check the details and confirm." + SIG,
           files=[("DEMO_SI.txt", si_txt()),
                  ("DEMO_BL.txt", bl_txt(with_(port_of_discharge="DAMMAM, SAUDI ARABIA (SADMM)", container_count="3 x 40'HC")))])),
    ("02-mismatch-excel-and-word.eml", "BL_COMPARISON", "MISMATCH", "consignee, gross_weight_kg",
     email("RE_ AFEMY - JEBEL ALI_UAE - MSC(MEDUDEMO0917) - 5RDM-00918 - INV77 - HARBOURLINE - CFR",
           "Dear Mitchelle,\n\nPlease compare the SI (Excel) and draft BL (Word) attached and revert with any discrepancy." + SIG,
           files=[("DEMO_SI.xlsx", si_xlsx()),
                  ("DEMO_BL.docx", bl_docx(with_(consignee="HARBOURLINE TRADING LLC", gross_weight_kg="87,400 KG")))])),
    ("03-clean-pair-all-match.eml", "BL_COMPARISON", "OK", "",
     email("REQUEST BL DRAFT _ PO 26917_ A4 COPIER PAPER__86MT",
           "Hi,\n\nPFA SI and draft BL for checking. Note the BL uses different labels (POD, To the Order of) for the same fields." + SIG,
           files=[("DEMO_SI.txt", si_txt()), ("DEMO_BL.txt", bl_txt())])),
    ("04-wrong-document-invoice-not-bl.eml", "BL_COMPARISON", "NEEDS_REVIEW", "wrong_doc_type",
     email("TO CONFIRM DOCS _ 5RDM-00919 _ JEBEL ALI_UAE _ HARBOURLINE",
           "Dear Team,\n\nPlease find attached the SI and the Commercial Invoice. Kindly confirm the BL is in order." + SIG,
           files=[("DEMO_SI.txt", si_txt()), ("DEMO_BL.txt", invoice_txt())])),
    ("05-missing-attachment.eml", "BL_COMPARISON", "NEEDS_REVIEW", "missing_attachment",
     email("TO CONFIRM DOCS _ 5RDM-00920 _ JEBEL ALI_UAE _ HARBOURLINE",
           "Dear Team,\n\nAttached SI and draft BL for checking, please compare the SI and draft BL and confirm." + SIG)),
    ("06-unreadable-empty-pdf.eml", "BL_COMPARISON", "NEEDS_REVIEW", "unreadable",
     email("TO CONFIRM DOCS _ 5RDM-00921 _ JEBEL ALI_UAE _ HARBOURLINE",
           "Dear Team,\n\nAttached SI and draft BL for checking (the BL file may not open). Please advise." + SIG,
           files=[("DEMO_SI.txt", si_txt()), ("DEMO_BL.pdf", b"")])),
    ("07-blank-field-in-si.eml", "BL_COMPARISON", "NEEDS_REVIEW", "missing_value",
     email("TO CONFIRM DOCS _ 5RDM-00922 _ JEBEL ALI_UAE _ HARBOURLINE",
           "Dear Team,\n\nPlease compare the SI and draft BL. The customer left some SI fields blank." + SIG,
           files=[("DEMO_SI.txt", si_txt(with_(notify_party="TBA", gross_weight_kg="???"))), ("DEMO_BL.txt", bl_txt())])),
    ("08-request-for-draft-no-files.eml", "BL_COMPARISON", "OK", "(nothing to compare yet)",
     email("Draft BL DEMO STAR V.026E PORT KLANG - amend BL 017",
           "Dear Hari,\n\nPlease assist to send the draft BL for DEMOBL260917 for checking asap." + SIG)),
    ("09-si-request.eml", "SI_REQUEST", "-", "",
     email("REQUEST SI _ 5RDM-00923 _ JEBEL ALI_UAE _ HARBOURLINE STATIONERY LLC _ DEMOBL260923",
           "Hi Willy,\n\nPlease find Shipping instruction for 5RDM-00923.\n\nPOL: PORT KLANG\nPOD: JEBEL ALI" + SIG)),
    ("10-invoice-query.eml", "INVOICE_QUERY", "-", "",
     email("REQUEST TO CANCEL INVOICE -5250091701 - HARBOURLINE - 5RDM-00924",
           "Hi,\n\nPlease cancel invoice 5250091701 (USD 1,250.00) and reissue under the correct consignee." + SIG,
           sender="finance@harbourline.example")),
    ("11-general-bot-notice.eml", "GENERAL", "-", "",
     email("_RPA_ India HSS SD Billing Process Completed - DEMO STAR V.026E",
           "This is an automated notification. The billing process for DEMO STAR V.026E has completed. No action required.",
           sender="rpa.bot@demotrading.example")),
    ("12-spam-phishing.eml", "SPAM", "-", "",
     email("URGENT: Your email storage is full - verify account immediately",
           "Dear user, your mailbox has exceeded its storage limit. Verify your account within 24 hours to avoid "
           "deactivation: http://webmail-verify.example/login", sender="alert@webmail-verify.co")),
    ("13-malay-email-mismatch.eml", "BL_COMPARISON", "MISMATCH", "shipper",
     email("TO CONFIRM DOCS _ 5RDM-00925 _ JEBEL ALI_UAE _ HARBOURLINE",
           "Salam sejahtera,\n\nBersama ini dilampirkan SI dan draf BL untuk OC 5RDM-00925. Sila semak butiran dan sahkan.\n"
           "Attached are the SI and draft BL.\n\nTerima kasih." + SIG,
           files=[("DEMO_SI.txt", si_txt()), ("DEMO_BL.txt", bl_txt(with_(shipper="DEMO PULP TRADING SDN BHD")))])),
]


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    for f in OUT.glob("*.eml"):
        f.unlink()
    for name, *_rest, raw in CASES:
        (OUT / name).write_bytes(raw)

    # Verify every case against the real pipeline.
    import tempfile
    from sdoc.config import settings
    from sdoc.inbox import open_inbox
    from sdoc.mail import CompositeInbox, UploadInbox
    from sdoc.pipeline import build_pipeline

    comp = CompositeInbox(open_inbox(settings.data_dir))
    up = UploadInbox(Path(tempfile.mkdtemp()))
    comp.add_source("upload", up)
    pipe = build_pipeline(settings, comp)
    ok = True
    print(f"{'file':44} {'category':14} {'status':13} detail")
    for name, cat, status, detail, raw in CASES:
        r = pipe.process(up.add(raw))
        got_status = r.status.value if r.category.value == "BL_COMPARISON" else "-"
        got_detail = r.review_reason.value if r.review_reason else ", ".join(sorted(r.defect_fields))
        good = r.category.value == cat and got_status == status and (not detail or detail.startswith("(") or got_detail == detail)
        ok &= good
        print(f"{'PASS' if good else 'FAIL'} {name:39} {r.category.value:14} {got_status:13} {got_detail}")
    print(f"\n{len(CASES)} files written to {OUT}")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
